"""Create a sample till export to try the report on.

The numbers are made up (a fixed random seed, so every run gives the same file),
but the file is shaped like a real export from an Israeli till: a title block
above the headers, dates stored sometimes as dates and sometimes as text, item
names typed slightly differently, refunds as negative quantities, blank rows and
a totals row at the bottom. Cleaning that up is half of what report.py does.

    python make_sample_data.py            -> data/till-export.xlsx
"""

import random
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

random.seed(5568)

CAFE = "קפה שדרה"
LAST_DAY = date(2026, 10, 10)  # a Saturday; the export covers the two weeks before it
DAYS = 14

# name, price, popularity weight, spellings the staff actually type
MENU = [
    ("אספרסו", 9, 10, ["אספרסו", "אספרסו ", "אקספרסו"]),
    ("אמריקנו", 12, 9, ["אמריקנו"]),
    ("הפוך", 14, 22, ["הפוך", "הפוך ", "קפה הפוך"]),
    ("קפה קר", 16, 7, ["קפה קר", "קפה קר "]),
    ("מאצ'ה לאטה", 18, 5, ["מאצ'ה לאטה", "מאצה לאטה"]),
    ("קרואסון חמאה", 16, 12, ["קרואסון חמאה", "קרואסון"]),
    ("בורקס גבינה", 14, 8, ["בורקס גבינה", "בורקס"]),
    ("ראגלך", 8, 6, ["ראגלך"]),
    ("עוגת שוקולד", 22, 5, ["עוגת שוקולד"]),
    ("שקשוקה", 48, 6, ["שקשוקה"]),
    ("טוסט גבינות", 32, 7, ["טוסט גבינות", "טוסט"]),
    ("סלט ירקות", 38, 4, ["סלט ירקות"]),
    ("כריך אבוקדו", 36, 5, ["כריך אבוקדו", "כריך אבוקדו "]),
]

# Opening hours of the demo café: Sunday to Thursday 7:30-19:00, Friday 7:30-15:00, closed Saturday.
HOURS = {6: (7, 19), 0: (7, 19), 1: (7, 19), 2: (7, 19), 3: (7, 19), 4: (7, 15)}  # date.weekday(): Mon=0 ... Sun=6

# How busy each hour is (morning rush, lunch, a small afternoon bump).
HOUR_WEIGHT = {7: 6, 8: 12, 9: 11, 10: 8, 11: 6, 12: 9, 13: 8, 14: 5, 15: 4, 16: 5, 17: 4, 18: 3}

PAYMENTS = [("אשראי", 70), ("ביט", 17), ("מזומן", 13)]


def pick(options):
    """Pick one option from (value, weight) pairs."""
    return random.choices([o[0] for o in options], weights=[o[1] for o in options])[0]


def main():
    rows = []
    receipt = 48210
    first = LAST_DAY - timedelta(days=DAYS - 1)
    for i in range(DAYS):
        day = first + timedelta(days=i)
        if day.weekday() not in HOURS:
            continue  # closed
        open_h, close_h = HOURS[day.weekday()]
        second_week = i >= 7
        busy = 1.06 if second_week else 1.0  # the second week is a bit busier
        if day.weekday() == 4:
            busy *= 1.35  # Fridays are short and packed
        receipts_today = int(random.gauss(118, 12) * busy)
        hours = [h for h in HOUR_WEIGHT if open_h <= h < close_h]
        for _ in range(receipts_today):
            receipt += 1
            hour = random.choices(hours, weights=[HOUR_WEIGHT[h] for h in hours])[0]
            minute = random.randint(30 if hour == 7 else 0, 59)
            payment = pick(PAYMENTS)
            for _ in range(random.choices([1, 2, 3, 4], weights=[45, 35, 15, 5])[0]):
                name, price, weight, spellings = random.choices(MENU, weights=[m[2] for m in MENU])[0]
                qty = 2 if random.random() < 0.08 else 1
                discount = round(price * qty * 0.1, 2) if random.random() < 0.05 else 0
                rows.append([day, time(hour, minute), receipt, spellings[0] if random.random() < 0.85 else random.choice(spellings), qty, price, discount, payment])
            # A hot second week: iced coffee sells more.
            if second_week and random.random() < 0.06:
                rows.append([day, time(hour, minute), receipt, "קפה קר", 1, 16, 0, payment])

    # A few refunds: same item, negative quantity, on a later receipt.
    for _ in range(9):
        src = random.choice(rows)
        receipt += 1
        rows.append([src[0], src[1], receipt, src[3], -src[4], src[5], 0, src[7]])
    rows.sort(key=lambda r: (r[0], r[1], r[2]))

    wb = Workbook()
    ws = wb.active
    ws.title = "דוח מכירות"
    ws.sheet_view.rightToLeft = True
    ws.append([f"קופה ראשית - {CAFE}"])
    ws.append([f"הופק: {LAST_DAY + timedelta(days=1):%d/%m/%Y} 06:58"])
    ws.append([])
    ws.append(["תאריך", "שעה", "מס' קבלה", "פריט", "כמות", "מחיר ליח'", "הנחה", "סה\"כ שורה", "אמצעי תשלום"])
    for c in ws[4]:
        c.font = Font(bold=True)

    total = 0
    for n, (day, t, rec, item, qty, price, discount, payment) in enumerate(rows):
        line = round(qty * price - discount, 2)
        total += line
        # Some tills write the date as text, and on some rows leave the line total empty.
        when = day.strftime("%d/%m/%Y") if n % 7 == 3 else datetime.combine(day, time())
        line_cell = None if n % 23 == 5 else line
        ws.append([when, t.strftime("%H:%M"), rec, item, qty, price, discount or None, line_cell, payment])
        if n % 150 == 77:
            ws.append([])  # stray blank rows
    ws.append([])
    ws.append(["סה\"כ", None, None, None, None, None, None, round(total, 2), None])

    for col, width in zip("ABCDEFGHI", [12, 8, 10, 16, 7, 9, 8, 10, 12]):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=5, max_col=1):
        row[0].number_format = "DD/MM/YYYY"

    out = Path(__file__).parent / "data" / "till-export.xlsx"
    out.parent.mkdir(exist_ok=True)
    wb.save(out)
    print(f"Wrote {out} ({len(rows)} item lines)")


if __name__ == "__main__":
    main()
