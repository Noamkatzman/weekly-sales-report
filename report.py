"""Weekly sales report: turns a till export into a one-page PDF and emails it.

    python report.py                          # last full week (Sunday to Saturday)
    python report.py --week-ending 2026-10-10 # a specific week
    python report.py --send                   # also email the PDF to the owner

Email settings come from environment variables (see README.md). Without them,
--send saves the email as an .eml file in the output folder instead, so you can
open it and see exactly what would have been sent.
"""

import argparse
import html
import os
import re
import smtplib
import subprocess
import sys
from collections import Counter, defaultdict
from datetime import date, datetime, time, timedelta
from email.message import EmailMessage
from pathlib import Path
from string import Template

from openpyxl import load_workbook

HERE = Path(__file__).parent

# ---- Settings for this business -------------------------------------------

CAFE_NAME = "קפה שדרה"
INPUT = HERE / "data" / "till-export.xlsx"
OUTPUT_DIR = HERE / "output"

# Spellings the staff type for the same item. Anything not listed is kept as is.
ALIASES = {
    "אקספרסו": "אספרסו",
    "קפה הפוך": "הפוך",
    "מאצה לאטה": "מאצ'ה לאטה",
    "קרואסון": "קרואסון חמאה",
    "בורקס": "בורקס גבינה",
    "טוסט": "טוסט גבינות",
}

# Column titles in the till export, mapped to the names used here.
COLUMNS = {
    "תאריך": "date",
    "שעה": "time",
    "מס' קבלה": "receipt",
    "פריט": "item",
    "כמות": "qty",
    "מחיר ליח'": "price",
    "הנחה": "discount",
    'סה"כ שורה': "total",
    "אמצעי תשלום": "payment",
}

EDGE_PATHS = [
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
]

DAY_NAMES = ["ראשון", "שני", "שלישי", "רביעי", "חמישי", "שישי", "שבת"]  # Sunday first
MONTHS = ["ינואר", "פברואר", "מרץ", "אפריל", "מאי", "יוני", "יולי", "אוגוסט", "ספטמבר", "אוקטובר", "נובמבר", "דצמבר"]


# ---- Reading and cleaning the export --------------------------------------

def read_export(path):
    """Return the item lines from the till export, cleaned, plus a count of what was fixed."""
    ws = load_workbook(path, read_only=True, data_only=True).active
    rows = ws.iter_rows(values_only=True)
    fixed = Counter()

    # Tills put a title block above the table, so look for the header row.
    for header in rows:
        if header and "תאריך" in header and "פריט" in header:
            break
    else:
        sys.exit(f"Couldn't find the header row (with 'תאריך' and 'פריט') in {path}")
    index = {COLUMNS[h]: i for i, h in enumerate(header) if h in COLUMNS}
    missing = set(COLUMNS.values()) - set(index) - {"discount", "total"}
    if missing:
        sys.exit(f"The export is missing these columns: {', '.join(sorted(missing))}")

    lines = []
    for row in rows:
        get = lambda key: row[index[key]] if key in index and index[key] < len(row) else None
        if not any(cell not in (None, "") for cell in row):
            fixed["blank rows skipped"] += 1
            continue
        if str(row[0] or "").startswith('סה"כ'):
            fixed["totals rows skipped"] += 1
            continue
        day = parse_date(get("date"))
        if day is None or get("receipt") is None:
            fixed["unreadable rows skipped"] += 1
            continue
        if not isinstance(get("date"), datetime):
            fixed["dates stored as text"] += 1

        raw_item = str(get("item") or "")
        item = re.sub(r"\s+", " ", raw_item).strip()
        item = ALIASES.get(item, item)
        if item != raw_item:
            fixed["item names merged"] += 1

        qty = float(get("qty") or 0)
        price = float(get("price") or 0)
        discount = float(get("discount") or 0)
        total = get("total")
        if total is None:
            total = qty * price - discount
            fixed["missing line totals filled"] += 1

        lines.append({
            "day": day,
            "hour": parse_hour(get("time")),
            "receipt": get("receipt"),
            "item": item,
            "qty": qty,
            "total": float(total),
            "payment": str(get("payment") or "לא ידוע").strip(),
        })
    return lines, fixed


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ("%d/%m/%Y", "%d.%m.%Y", "%d/%m/%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    return None


def parse_hour(value):
    if isinstance(value, (time, datetime)):
        return value.hour
    match = re.match(r"\s*(\d{1,2})[:.]", str(value or ""))
    return int(match.group(1)) if match else None


# ---- The numbers ----------------------------------------------------------

def summarize(lines, start):
    """Totals for the 7 days starting on `start` (a Sunday)."""
    week = [l for l in lines if start <= l["day"] < start + timedelta(days=7)]
    sales = [l for l in week if l["qty"] > 0]
    refunds = [l for l in week if l["qty"] < 0]

    receipts = {l["receipt"] for l in sales}
    revenue = sum(l["total"] for l in week)  # refunds are negative, so this is net
    by_day = [0.0] * 7
    for l in week:
        by_day[(l["day"] - start).days] += l["total"]

    receipts_by_hour = defaultdict(set)
    for l in sales:
        if l["hour"] is not None:
            receipts_by_hour[l["hour"]].add(l["receipt"])

    items = defaultdict(lambda: {"qty": 0.0, "revenue": 0.0})
    for l in sales:
        items[l["item"]]["qty"] += l["qty"]
        items[l["item"]]["revenue"] += l["total"]

    payments = Counter()
    for l in sales:
        payments[l["payment"]] += l["total"]

    return {
        "start": start,
        "revenue": revenue,
        "receipts": len(receipts),
        "average": sum(l["total"] for l in sales) / len(receipts) if receipts else 0,
        "refund_count": len(refunds),
        "refund_total": -sum(l["total"] for l in refunds),
        "by_day": by_day,
        "by_hour": {h: len(r) for h, r in sorted(receipts_by_hour.items())},
        "items": dict(items),
        "payments": payments,
        "open_days": sum(1 for v in by_day if v > 0),
    }


def change(now, before):
    return None if not before else (now - before) / before * 100


def insights(this, last):
    """A few plain sentences an owner can read in ten seconds."""
    out = []
    best = max(range(7), key=lambda d: this["by_day"][d])
    out.append(f"היום החזק בשבוע היה יום {DAY_NAMES[best]}, עם {shekels(this['by_day'][best])}.")

    if this["by_hour"]:
        hour = max(this["by_hour"], key=this["by_hour"].get)
        per_day = this["by_hour"][hour] / max(this["open_days"], 1)
        out.append(f"השעה העמוסה: {ltr(f'{hour}:00–{hour + 1}:00')}, כ-{per_day:.0f} קבלות ביום. כדאי שיהיו שני אנשים מאחורי הדלפק.")

    movers = []
    for name, now in this["items"].items():
        before = last["items"].get(name, {}).get("qty", 0)
        if before >= 15:
            movers.append((change(now["qty"], before), name))
    if movers:
        pct, name = max(movers, key=lambda m: abs(m[0]))
        word = "עלו" if pct > 0 else "ירדו"
        out.append(f"המכירות של {name} {word} ב-{abs(pct):.0f}% לעומת השבוע הקודם.")

    if this["refund_count"]:
        out.append(f"{this['refund_count']} החזרים בסך {shekels(this['refund_total'])}.")
    return out


# ---- The report -----------------------------------------------------------

def ltr(text):
    """Keep a run like "4–10" or "9:00–10:00" left to right inside Hebrew text."""
    return f"⁦{text}⁩"


def shekels(value):
    return f"₪{value:,.0f}"


def pct_badge(pct, invert=False):
    if pct is None:
        return '<span class="chg">—</span>'
    good = (pct >= 0) != invert
    arrow = "▲" if pct >= 0 else "▼"
    return f'<span class="chg {"up" if good else "down"}"><bdi dir="ltr">{arrow} {abs(pct):.1f}%</bdi></span>'


def date_range(start):
    end = start + timedelta(days=6)
    if start.month == end.month:
        return f"{ltr(f'{start.day}–{end.day}')} ב{MONTHS[end.month - 1]} {end.year}"
    return f"{start.day} ב{MONTHS[start.month - 1]} – {end.day} ב{MONTHS[end.month - 1]} {end.year}"


def day_chart(this, last):
    """Bars per day: last week (light) behind this week (dark)."""
    w, h, top, bottom = 640, 220, 14, 30
    peak = max(this["by_day"] + last["by_day"]) or 1
    col = w / 7
    parts = []
    for d in range(7):
        # Sunday on the right, like a Hebrew calendar.
        x = w - (d + 1) * col
        for value, offset, cls in ((last["by_day"][d], 0.22, "prev"), (this["by_day"][d], 0.48, "now")):
            bh = (h - top - bottom) * value / peak
            parts.append(f'<rect class="{cls}" x="{x + col * offset:.1f}" y="{h - bottom - bh:.1f}" width="{col * 0.3:.1f}" height="{bh:.1f}" rx="3"/>')
        label = DAY_NAMES[d] if this["by_day"][d] or last["by_day"][d] else f"{DAY_NAMES[d]} (סגור)"
        parts.append(f'<text x="{x + col / 2:.1f}" y="{h - 9}" text-anchor="middle">{label}</text>')
        if this["by_day"][d]:
            parts.append(f'<text class="val" x="{x + col * 0.63:.1f}" y="{h - bottom - (h - top - bottom) * this["by_day"][d] / peak - 5:.1f}" text-anchor="middle">{this["by_day"][d] / 1000:.1f}K</text>')
    return f'<svg viewBox="0 0 {w} {h}" role="img" aria-label="הכנסות לפי יום">{"".join(parts)}</svg>'


def hour_chart(this):
    hours = this["by_hour"]
    if not hours:
        return ""
    w, h, bottom = 300, 130, 22
    peak = max(hours.values())
    keys = list(range(min(hours), max(hours) + 1))
    col = w / len(keys)
    parts = []
    for i, hour in enumerate(keys):
        value = hours.get(hour, 0)
        x = w - (i + 1) * col
        bh = (h - bottom - 8) * value / peak
        cls = "now" if value == peak else "prev"
        parts.append(f'<rect class="{cls}" x="{x + col * 0.15:.1f}" y="{h - bottom - bh:.1f}" width="{col * 0.7:.1f}" height="{bh:.1f}" rx="2"/>')
        parts.append(f'<text x="{x + col / 2:.1f}" y="{h - 6}" text-anchor="middle">{hour}</text>')
    return f'<svg viewBox="0 0 {w} {h}" role="img" aria-label="קבלות לפי שעה">{"".join(parts)}</svg>'


def items_table(this, last):
    top = sorted(this["items"].items(), key=lambda kv: kv[1]["revenue"], reverse=True)[:6]
    peak = top[0][1]["revenue"] if top else 1
    rows = []
    for rank, (name, v) in enumerate(top, 1):
        before = last["items"].get(name, {}).get("qty")
        rows.append(
            f'<tr><td class="rank">{rank}</td>'
            f'<td><div class="item">{html.escape(name)}</div><div class="bar" style="width:{v["revenue"] / peak * 100:.0f}%"></div></td>'
            f'<td class="num">{v["qty"]:.0f}</td><td class="num">{shekels(v["revenue"])}</td>'
            f'<td class="num">{pct_badge(change(v["qty"], before))}</td></tr>'
        )
    return "".join(rows)


def payments_list(this):
    total = sum(this["payments"].values()) or 1
    out = []
    for name, value in this["payments"].most_common():
        share = value / total * 100
        out.append(
            f'<li><span>{html.escape(name)}</span><span class="track"><span style="width:{share:.0f}%"></span></span>'
            f'<span class="num">{share:.0f}%</span></li>'
        )
    return "".join(out)


def render(this, last, fixed, line_count, source):
    template = Template((HERE / "template.html").read_text(encoding="utf-8"))
    return template.substitute(
        cafe=html.escape(CAFE_NAME),
        range=date_range(this["start"]),
        prev_range=date_range(last["start"]),
        revenue=shekels(this["revenue"]),
        revenue_chg=pct_badge(change(this["revenue"], last["revenue"])),
        receipts=f"{this['receipts']:,}",
        receipts_chg=pct_badge(change(this["receipts"], last["receipts"])),
        average=shekels(this["average"]),
        average_chg=pct_badge(change(this["average"], last["average"])),
        refunds=shekels(this["refund_total"]),
        refunds_note=f"{this['refund_count']} החזרים",
        insights="".join(f"<li>{html.escape(s)}</li>" for s in insights(this, last)),
        day_chart=day_chart(this, last),
        hour_chart=hour_chart(this),
        items=items_table(this, last),
        payments=payments_list(this),
        footer=html.escape(
            f"נוצר אוטומטית מ-{source.name} ב-{datetime.now():%d/%m/%Y %H:%M} · "
            f"{line_count:,} שורות נקראו, {sum(fixed.values())} תוקנו או דולגו"
        ),
    )


def to_pdf(html_path, pdf_path):
    browser = next((p for p in EDGE_PATHS if Path(p).exists()), None)
    if not browser:
        print("  No Edge or Chrome found, so no PDF. The HTML report is ready.")
        return False
    subprocess.run(
        [browser, "--headless", "--disable-gpu", "--no-pdf-header-footer",
         f"--print-to-pdf={pdf_path}", html_path.resolve().as_uri()],
        check=True, capture_output=True, timeout=60,
    )
    return pdf_path.exists()


# ---- Email ----------------------------------------------------------------

def build_email(this, last, pdf_path):
    msg = EmailMessage()
    msg["Subject"] = f"דוח מכירות שבועי – {CAFE_NAME} – {date_range(this['start'])}"
    msg["From"] = os.environ.get("REPORT_FROM") or os.environ.get("SMTP_USER") or "reports@example.com"
    msg["To"] = os.environ.get("REPORT_TO", "owner@example.com")
    rev = change(this["revenue"], last["revenue"])
    body = [
        "בוקר טוב,",
        "",
        f"הדוח של השבוע שעבר ({date_range(this['start'])}) מצורף כ-PDF. בקצרה:",
        f"• הכנסות: {shekels(this['revenue'])}" + (f" ({rev:+.1f}% לעומת השבוע הקודם)" if rev is not None else ""),
        f"• {this['receipts']:,} קבלות, קבלה ממוצעת {shekels(this['average'])}",
        *[f"• {s}" for s in insights(this, last)],
        "",
        "הדוח הזה נשלח אוטומטית כל יום ראשון.",
    ]
    msg.set_content("\n".join(body))
    msg.add_attachment(pdf_path.read_bytes(), maintype="application", subtype="pdf", filename=pdf_path.name)
    return msg


def send(msg, eml_path):
    host = os.environ.get("SMTP_HOST")
    if not host:
        eml_path.write_bytes(bytes(msg))
        print(f"  SMTP_HOST isn't set, so the email was saved instead: {eml_path.name}")
        return
    port = int(os.environ.get("SMTP_PORT", "587"))
    with smtplib.SMTP(host, port, timeout=30) as smtp:
        if port != 25:
            smtp.starttls()
        if os.environ.get("SMTP_USER"):
            smtp.login(os.environ["SMTP_USER"], os.environ["SMTP_PASSWORD"])
        smtp.send_message(msg)
    print(f"  Sent to {msg['To']}")


# ---- Run ------------------------------------------------------------------

def last_saturday(today):
    """The most recent Saturday before today (on a Saturday, the week isn't over yet)."""
    return today - timedelta(days=(today.weekday() - 5) % 7 or 7)


def main():
    parser = argparse.ArgumentParser(description="Weekly sales report from a till export.")
    parser.add_argument("--input", type=Path, default=INPUT)
    parser.add_argument("--week-ending", type=date.fromisoformat, help="the Saturday that ends the week (YYYY-MM-DD)")
    parser.add_argument("--send", action="store_true", help="email the PDF")
    args = parser.parse_args()

    week_end = args.week_ending or last_saturday(date.today())
    if week_end.weekday() != 5:
        sys.exit(f"--week-ending should be a Saturday; {week_end} is not.")
    start = week_end - timedelta(days=6)

    print(f"Reading {args.input.name}")
    lines, fixed = read_export(args.input)
    print(f"  {len(lines):,} item lines from {min(l['day'] for l in lines):%d/%m} to {max(l['day'] for l in lines):%d/%m}")
    for what, n in fixed.most_common():
        print(f"  {n:>4}  {what}")

    this = summarize(lines, start)
    last = summarize(lines, start - timedelta(days=7))
    if not this["receipts"]:
        sys.exit(f"No sales in the week {start} to {week_end}. Is this the right export?")
    rev = change(this["revenue"], last["revenue"])
    print(f"Week {start:%d/%m}-{week_end:%d/%m}: revenue {this['revenue']:,.0f} ILS"
          + (f" ({rev:+.1f}%)" if rev is not None else "") + f", {this['receipts']:,} receipts")

    OUTPUT_DIR.mkdir(exist_ok=True)
    name = f"sales-report-{week_end:%Y-%m-%d}"
    html_path = OUTPUT_DIR / f"{name}.html"
    html_path.write_text(render(this, last, fixed, len(lines), args.input), encoding="utf-8")
    pdf_path = OUTPUT_DIR / f"{name}.pdf"
    if to_pdf(html_path, pdf_path):
        print(f"Report: output/{pdf_path.name}")

    if args.send:
        if not pdf_path.exists():
            sys.exit("No PDF to send.")
        send(build_email(this, last, pdf_path), OUTPUT_DIR / f"{name}.eml")


if __name__ == "__main__":
    main()
